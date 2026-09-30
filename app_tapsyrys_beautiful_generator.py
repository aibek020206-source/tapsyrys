import base64
import hashlib
import html
import io
import json
import os
import re
import secrets
import sqlite3
import time
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

from flask import Flask, redirect, request, make_response, send_file

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY") or secrets.token_hex(32)

# =========================
# CANVA OAUTH
# =========================
CLIENT_ID = os.environ.get("CANVA_CLIENT_ID")
CLIENT_SECRET = os.environ.get("CANVA_CLIENT_SECRET")

REDIRECT_URI = os.environ.get(
    "CANVA_REDIRECT_URI",
    "https://tapsyrys.onrender.com/canva/callback",
)

CANVA_AUTHORIZE_URL = "https://www.canva.com/api/oauth/authorize"
CANVA_TOKEN_URL = "https://api.canva.com/rest/v1/oauth/token"

SCOPES = "design:content:read design:meta:read profile:read"

STATE_TTL = 600
DB_PATH = os.environ.get("DB_PATH", "app.db")

# =========================
# APP SETTINGS
# =========================
SERVICES = {
    "Презентация": {"type": "page", "price": 60},
}

MAX_SLIDES = 30

# Wikimedia Commons API:
# API key қажет емес.
WIKIMEDIA_API = "https://commons.wikimedia.org/w/api.php"
USER_AGENT = "TapsyrysAI/1.0 (educational presentation generator)"


# =========================
# DATABASE
# =========================
def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    with db() as conn:
        conn.execute("""
            CREATE TABLE IF NOT EXISTS oauth_states (
                state TEXT PRIMARY KEY,
                code_verifier TEXT NOT NULL,
                created_at REAL NOT NULL
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS canva_tokens (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                access_token TEXT NOT NULL,
                refresh_token TEXT,
                expires_at REAL,
                scope TEXT,
                created_at REAL NOT NULL
            )
        """)

        conn.execute("""
            CREATE TABLE IF NOT EXISTS customers (
                customer_id TEXT PRIMARY KEY,
                free_used INTEGER NOT NULL DEFAULT 0,
                presentation_free_used INTEGER NOT NULL DEFAULT 0,
                created_at REAL NOT NULL
            )
        """)

        try:
            conn.execute(
                "ALTER TABLE customers "
                "ADD COLUMN presentation_free_used INTEGER NOT NULL DEFAULT 0"
            )
        except sqlite3.OperationalError:
            pass

        conn.execute("""
            CREATE TABLE IF NOT EXISTS orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                customer_id TEXT NOT NULL,
                service TEXT NOT NULL,
                subject TEXT,
                topic TEXT NOT NULL,
                quantity INTEGER NOT NULL DEFAULT 1,
                due_date TEXT,
                due_time TEXT,
                price INTEGER NOT NULL,
                payment_status TEXT NOT NULL DEFAULT 'pending',
                generation_status TEXT NOT NULL DEFAULT 'locked',
                created_at REAL NOT NULL
            )
        """)


def save_state(state, code_verifier):
    with db() as conn:
        conn.execute(
            "DELETE FROM oauth_states WHERE created_at < ?",
            (time.time() - STATE_TTL,),
        )
        conn.execute(
            """
            INSERT INTO oauth_states
            (state, code_verifier, created_at)
            VALUES (?, ?, ?)
            """,
            (state, code_verifier, time.time()),
        )


def pop_state(state):
    with db() as conn:
        row = conn.execute(
            "SELECT code_verifier, created_at "
            "FROM oauth_states WHERE state = ?",
            (state,),
        ).fetchone()

        conn.execute(
            "DELETE FROM oauth_states WHERE state = ?",
            (state,),
        )

    if not row or time.time() - row["created_at"] > STATE_TTL:
        return None

    return row["code_verifier"]


def save_tokens(token_data):
    expires_in = token_data.get("expires_in")
    expires_at = time.time() + expires_in if expires_in else None

    with db() as conn:
        conn.execute(
            """
            INSERT INTO canva_tokens
            (access_token, refresh_token, expires_at, scope, created_at)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                token_data["access_token"],
                token_data.get("refresh_token"),
                expires_at,
                token_data.get("scope"),
                time.time(),
            ),
        )


def get_customer(customer_id):
    with db() as conn:
        return conn.execute(
            "SELECT * FROM customers WHERE customer_id = ?",
            (customer_id,),
        ).fetchone()


def ensure_customer(customer_id):
    with db() as conn:
        conn.execute(
            """
            INSERT OR IGNORE INTO customers
            (customer_id, free_used, presentation_free_used, created_at)
            VALUES (?, 0, 0, ?)
            """,
            (customer_id, time.time()),
        )


def create_customer():
    return secrets.token_urlsafe(24)


def calculate_price(service, quantity):
    item = SERVICES[service]

    if item["type"] == "fixed":
        return item["price"]

    return item["price"] * max(1, quantity)


def customer_id_from_request():
    return request.cookies.get("tapsyrys_customer_id")


init_db()


# =========================
# LOCAL PRESENTATION CONTENT
# =========================
def make_local_slides(topic, subject, quantity):
    """
    OpenAI API қолданбай презентация мазмұнын құрады.
    """

    quantity = max(1, min(int(quantity), MAX_SLIDES))

    topic = re.sub(
        r"\s+",
        " ",
        (topic or "Тақырып").strip()
    )

    subject = re.sub(
        r"\s+",
        " ",
        (subject or "Пән").strip()
    )

    stop_words = {
        "және", "мен", "үшін", "туралы", "бойынша",
        "негіздері", "негізі", "кіріспе", "тақырыбы",
        "пәні", "пән", "the", "and", "of", "in",
        "на", "и", "для", "по", "о"
    }

    words = re.findall(
        r"[A-Za-zА-Яа-яӘәҒғҚқҢңӨөҰұҮүҺһІі0-9-]{3,}",
        topic.lower()
    )

    keywords = [
        word for word in words
        if word not in stop_words
    ]

    if not keywords:
        keywords = [topic]

    main_keyword = keywords[0]
    keyword_text = ", ".join(keywords[:4])

    subject_context = {
        "Философия":
            "ұғымдардың мәнін, себеп-салдарын және дүниетанымдық маңызын",
        "Психология":
            "адамның мінез-құлқы, ойлауы және эмоциялық ерекшеліктері тұрғысынан",
        "Социология":
            "қоғам, әлеуметтік топтар және қоғамдық байланыстар тұрғысынан",
        "Саясаттану":
            "саяси үдерістер, институттар және қоғамдық қатынастар тұрғысынан",
        "Тарих":
            "тарихи кезеңдер, оқиғалар және олардың себеп-салдары тұрғысынан",
        "Экология":
            "қоршаған орта, табиғи ресурстар және тұрақты даму тұрғысынан",
        "Физика":
            "физикалық құбылыстар, заңдылықтар және практикалық қолданылуы тұрғысынан",
        "Информатика":
            "ақпараттық технологиялар, алгоритмдер және цифрлық шешімдер тұрғысынан",
    }

    context = subject_context.get(
        subject,
        f"{subject} пәні аясында"
    )

    sections = [
        (
            "Кіріспе",
            [
                f"{topic} — {subject} пәніндегі маңызды тақырыптардың бірі.",
                f"Бұл презентацияда {main_keyword} ұғымы және оның негізгі қырлары қарастырылады.",
                "Негізгі мақсат: тақырыпты түсінікті түрде ашып көрсету.",
            ],
        ),
        (
            "Негізгі ұғымдар",
            [
                f"{topic} бойынша негізгі түсініктер мен терминдер қарастырылады.",
                f"Негізгі сөздер: {keyword_text}.",
                "Терминдерді дұрыс түсіну тақырыпты әрі қарай талдауға мүмкіндік береді.",
            ],
        ),
        (
            "Тақырыптың мақсаты",
            [
                f"{topic} мазмұнын жүйелі түрде түсіндіру.",
                f"Негізгі мәселелерді {context} талдау.",
                "Теориялық ақпаратты практикалық мысалдармен байланыстыру.",
            ],
        ),
        (
            "Тарихы және қалыптасуы",
            [
                f"{topic} түсінігінің қалыптасуына әртүрлі ғылыми және қоғамдық факторлар әсер етеді.",
                "Уақыт өте келе тақырыптың мазмұны кеңейіп, жаңа бағыттар пайда болды.",
                "Қазіргі түсінікті қалыптастыруда зерттеулердің маңызы жоғары.",
            ],
        ),
        (
            "Негізгі ерекшеліктері",
            [
                f"{topic} тақырыбының басты ерекшеліктері оның құрылымы мен мазмұнынан көрінеді.",
                f"Маңызды аспектілер: {keyword_text}.",
                "Әр ерекшелік тақырыпты толық түсінуге көмектеседі.",
            ],
        ),
        (
            "Құрылымы",
            [
                f"{topic} бірнеше өзара байланысты элементтен тұрады.",
                "Элементтер бір-біріне әсер етіп, жалпы жүйені қалыптастырады.",
                "Құрылымдық байланыстарды түсіну практикалық талдау үшін маңызды.",
            ],
        ),
        (
            "Негізгі мәселелер",
            [
                f"{topic} бойынша шешімін қажет ететін бірнеше мәселе бар.",
                "Мәселелердің пайда болуына әлеуметтік, ғылыми немесе практикалық факторлар әсер етуі мүмкін.",
                "Оларды шешу үшін нақты жағдайды жан-жақты талдау қажет.",
            ],
        ),
        (
            "Артықшылықтары",
            [
                f"{topic} қолданылуы немесе зерттелуі бірқатар пайдалы нәтиже береді.",
                "Негізгі артықшылықтары тиімділік, түсініктілік және практикалық маңызбен байланысты.",
                "Артықшылықтар нақты жағдайға қарай өзгеруі мүмкін.",
            ],
        ),
        (
            "Кемшіліктері мен шектеулері",
            [
                "Кез келген бағыттың белгілі бір шектеулері болады.",
                f"{topic} бойынша ақпаратты қолданғанда жағдайдың ерекшеліктерін ескеру қажет.",
                "Шектеулерді дұрыс бағалау қате қорытынды жасаудан сақтайды.",
            ],
        ),
        (
            "Практикалық қолданылуы",
            [
                f"{topic} теориямен ғана шектелмей, практикада да қолданылуы мүмкін.",
                f"Оны {context} қарастыруға болады.",
                "Практикалық мысалдар тақырыпты жақсы түсінуге мүмкіндік береді.",
            ],
        ),
        (
            "Мысал",
            [
                f"{topic} бойынша қарапайым мысал ретінде күнделікті өмірдегі нақты жағдайды алуға болады.",
                f"Мысалда {main_keyword} қалай көрінетіні көрсетіледі.",
                "Мысалды талдау теориялық түсінікті бекітеді.",
            ],
        ),
        (
            "Қазіргі кездегі маңызы",
            [
                f"Бүгінгі таңда {topic} тақырыбының өзектілігі сақталып отыр.",
                "Жаңа технологиялар мен қоғамдық өзгерістер тақырыпқа жаңа талаптар қояды.",
                "Сондықтан мәселені қазіргі жағдаймен байланыстырып қарастыру маңызды.",
            ],
        ),
        (
            "Қорытынды",
            [
                f"{topic} — зерттеуге және түсінуге маңызды тақырып.",
                "Негізгі ұғымдар мен ерекшеліктерді жүйелеу тақырыптың мәнін ашады.",
                "Тақырыпты практикамен байланыстыру оның маңызын нақты көрсетеді.",
            ],
        ),
        (
            "Пайдаланылған әдебиеттер",
            [
                "Пән бойынша оқу құралдары мен дәріс материалдары.",
                "Ғылыми мақалалар мен оқу әдебиеттері.",
                "Ресми және электрондық ақпараттық ресурстар.",
            ],
        ),
    ]

    if quantity == 1:
        selected = [
            (
                "Презентация",
                [
                    f"Тақырып: {topic}",
                    f"Пән: {subject}",
                    "Негізгі мазмұн қысқаша берілді.",
                ],
            )
        ]

    elif quantity < len(sections):
        selected = sections[:quantity]

        if quantity >= 3:
            selected[-1] = sections[-2]

    else:
        selected = list(sections)
        extra_index = 1

        while len(selected) < quantity:
            selected.insert(
                -1,
                (
                    f"Қосымша талдау {extra_index}",
                    [
                        f"{topic} тақырыбындағы {main_keyword} ұғымына қосымша талдау.",
                        f"Бұл бөлім {subject} пәнімен байланысты мәселелерді нақтылайды.",
                        "Нақты мысалдар мен деректерді қосу презентацияны толықтырады.",
                    ],
                ),
            )
            extra_index += 1

    return {
        "title": topic,
        "subject": subject,
        "slides": [
            {
                "title": title,
                "bullets": bullets,
            }
            for title, bullets in selected
        ],
    }


# =========================
# WIKIMEDIA COMMONS IMAGES
# =========================
def clean_search_text(text):
    text = re.sub(r"[^\w\sӘәҒғҚқҢңӨөҰұҮүҺһІі-]", " ", text or "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def search_wikimedia_image(query):
    """
    Wikimedia Commons-тен тегін сурет іздейді.
    API key қажет емес.

    Әр сурет үшін:
    - image URL
    - page URL
    - title
    қайтарылады.
    """

    query = clean_search_text(query)

    if not query:
        return None

    params = {
        "action": "query",
        "generator": "search",
        "gsrsearch": query,
        "gsrnamespace": "6",
        "gsrlimit": "20",
        "prop": "imageinfo",
        "iiprop": "url",
        "iiurlwidth": "1200",
        "format": "json",
        "formatversion": "2",
    }

    url = WIKIMEDIA_API + "?" + urllib.parse.urlencode(params)

    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/json",
        },
    )

    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            data = json.loads(
                response.read().decode("utf-8")
            )

        pages = data.get("query", {}).get("pages", [])

        for page in pages:
            imageinfo = page.get("imageinfo", [])

            if not imageinfo:
                continue

            info = imageinfo[0]

            image_url = (
                info.get("thumburl")
                or info.get("url")
            )

            if not image_url:
                continue

            title = page.get("title", "Wikimedia Commons")

            return {
                "image_url": image_url,
                "page_url": (
                    "https://commons.wikimedia.org/wiki/"
                    + urllib.parse.quote(
                        title.replace(" ", "_"),
                        safe="_:/"
                    )
                ),
                "title": title,
            }

    except Exception:
        return None

    return None


def download_image(image_url):
    """
    Суретті жадыға жүктейді.
    Дискке уақытша сурет сақтамайды.
    """

    req = urllib.request.Request(
        image_url,
        headers={
            "User-Agent": USER_AGENT,
        },
    )

    try:
        with urllib.request.urlopen(req, timeout=20) as response:
            data = response.read()

        if not data or len(data) < 1000:
            return None

        return io.BytesIO(data)

    except Exception:
        return None


def find_slide_image(topic, subject, slide_title, extra_text="", used_urls=None):
    """
    Слайдтың нақты мазмұнына сәйкес сурет іздейді.
    Әр слайдқа мүмкіндігінше бөлек сурет таңдалады.
    """

    used_urls = used_urls if used_urls is not None else set()

    topic = clean_search_text(topic)
    subject = clean_search_text(subject)
    slide_title = clean_search_text(slide_title)
    extra_text = clean_search_text(extra_text)

    # Ең нақты іздеу бірінші орындалады.
    queries = [
        f"{topic} {slide_title} {extra_text}".strip(),
        f"{topic} {slide_title}".strip(),
        f"{subject} {slide_title}".strip(),
        f"{topic}".strip(),
        f"{subject}".strip(),
    ]

    used_queries = set()

    for query in queries:
        query = clean_search_text(query)
        if not query or query.lower() in used_queries:
            continue
        used_queries.add(query.lower())

        results = search_wikimedia_images(query)
        for result in results:
            image_url = result.get("image_url")
            if not image_url or image_url in used_urls:
                continue

            image_stream = download_image(image_url)
            if image_stream:
                used_urls.add(image_url)
                return {
                    **result,
                    "stream": image_stream,
                }

    return None


def search_wikimedia_images(query):
    """Wikimedia Commons-тен бірнеше ықтимал сурет қайтарады."""
    query = clean_search_text(query)
    if not query:
        return []

    params = {
        "action": "query",
        "generator": "search",
        "gsrsearch": query,
        "gsrnamespace": "6",
        "gsrlimit": "20",
        "prop": "imageinfo",
        "iiprop": "url|mime",
        "iiurlwidth": "1200",
        "format": "json",
        "formatversion": "2",
    }

    url = WIKIMEDIA_API + "?" + urllib.parse.urlencode(params)
    req = urllib.request.Request(
        url,
        headers={
            "User-Agent": USER_AGENT,
            "Accept": "application/json",
        },
    )

    try:
        with urllib.request.urlopen(req, timeout=15) as response:
            data = json.loads(response.read().decode("utf-8"))

        pages = data.get("query", {}).get("pages", [])
        results = []

        for page in pages:
            imageinfo = page.get("imageinfo", [])
            if not imageinfo:
                continue

            info = imageinfo[0]
            mime = info.get("mime", "")
            if not mime.startswith("image/"):
                continue

            image_url = info.get("thumburl") or info.get("url")
            if not image_url:
                continue

            title = page.get("title", "Wikimedia Commons")
            results.append({
                "image_url": image_url,
                "page_url": (
                    "https://commons.wikimedia.org/wiki/"
                    + urllib.parse.quote(title.replace(" ", "_"), safe="_:/")
                ),
                "title": title,
            })

        return results
    except Exception:
        return []


def search_wikimedia_image(query):
    """Бір сурет керек болған ескі шақырулар үшін үйлесімді функция."""
    results = search_wikimedia_images(query)
    return results[0] if results else None


# =========================
# POWERPOINT
# =========================
def add_text_box(slide, text, left, top, width, height,
                 font_size=18, bold=False, color=(25, 32, 50),
                 font_name="Aptos", align=None):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    box = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)

    p = tf.paragraphs[0]
    p.text = str(text)
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.name = font_name
    p.font.color.rgb = RGBColor(*color)
    if align is not None:
        p.alignment = align
    return box


def add_image_safely(slide, image_stream, left, top, width, height):
    """
    Суретті слайдқа қосады.
    Сурет қосылмаса, презентация тоқтамайды.
    """

    from pptx.util import Inches

    try:
        image_stream.seek(0)

        slide.shapes.add_picture(
            image_stream,
            Inches(left),
            Inches(top),
            width=Inches(width),
            height=Inches(height),
        )

        return True

    except Exception:
        return False


def make_pptx(slides_data, order_id):
    """Заманауи, әр слайдта әртүрлі layout қолданылатын PPTX генераторы."""
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise RuntimeError("python-pptx орнатылмаған.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title = str(slides_data.get("title", "Презентация"))
    subject = str(slides_data.get("subject", ""))
    slides = slides_data.get("slides", [])
    used = set()

    # Біртұтас заманауи палитра
    NAVY = (18, 25, 43)
    INK = (29, 36, 55)
    MUTED = (95, 105, 125)
    WHITE = (255, 255, 255)
    LIGHT = (246, 248, 252)
    ACCENTS = [(76, 93, 220), (24, 150, 137), (236, 126, 70), (157, 78, 221)]

    def rect(slide, x, y, w, h, color, radius=False, line=None):
        shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb = RGBColor(*color)
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = RGBColor(*line)
        return s

    def add_bullets(slide, bullets, x, y, w, h, size=19, color=INK):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame; tf.word_wrap = True
        tf.margin_left = Inches(.03); tf.margin_right = Inches(.03)
        for i, bullet in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "•  " + str(bullet)
            p.font.name = "Aptos"; p.font.size = Pt(size); p.font.color.rgb = RGBColor(*color)
            p.space_after = Pt(12)
        return box

    def add_number(slide, n, color):
        add_text_box(slide, f"{n:02d}", .65, 6.95, .7, .25, 9, True, color)
        add_text_box(slide, "TAPSYRYS AI", 10.85, 6.95, 1.8, .25, 8, True, MUTED, align=PP_ALIGN.RIGHT)

    # 1. Мұқаба — үлкен сурет + қараңғы панель
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rect(slide, 0, 0, 13.333, 7.5, NAVY)
    cover = find_slide_image(title, subject, "мұқаба", used_urls=used)
    if cover:
        add_image_safely(slide, cover["stream"], 7.55, 0, 5.783, 7.5)
        rect(slide, 7.55, 0, 5.783, 7.5, (10, 15, 28))
        # Жеңіл overlay орнына сол жақ панельді анық қалдырамыз
        add_image_safely(slide, cover["stream"], 7.75, .35, 5.38, 6.8)
    rect(slide, .65, .7, .9, .08, ACCENTS[0])
    add_text_box(slide, subject.upper() if subject else "ПРЕЗЕНТАЦИЯ", .65, 1.1, 5.9, .45, 13, True, (160, 170, 195))
    add_text_box(slide, title, .65, 2.0, 6.4, 2.0, 34, True, WHITE)
    add_text_box(slide, "Қысқа, түсінікті және заманауи презентация", .65, 4.45, 5.8, .65, 16, False, (205, 211, 225))
    add_text_box(slide, "Tapsyrys AI", .65, 6.75, 2.0, .35, 11, True, WHITE)
    if cover:
        add_text_box(slide, "Wikimedia Commons", 9.0, 6.95, 3.7, .25, 8, False, WHITE, align=PP_ALIGN.RIGHT)

    # 2+. Әртүрлі layout
    for idx, item in enumerate(slides, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        accent = ACCENTS[(idx - 1) % len(ACCENTS)]
        stitle = str(item.get("title", ""))
        bullets = [str(x) for x in item.get("bullets", []) if str(x).strip()]
        image = find_slide_image(title, subject, stitle, " ".join(bullets[:2]), used_urls=used)
        layout_no = (idx - 1) % 4

        if layout_no == 0:
            # SPLIT — мәтін + үлкен сурет
            rect(slide, 0, 0, 13.333, 7.5, LIGHT)
            rect(slide, 0, 0, .16, 7.5, accent)
            add_text_box(slide, stitle, .7, .65, 7.0, 1.0, 27, True, INK)
            add_bullets(slide, bullets, .75, 1.8, 6.5, 4.5, 18)
            if image:
                rect(slide, 8.0, 1.25, 4.55, 5.35, WHITE, True)
                add_image_safely(slide, image["stream"], 8.15, 1.4, 4.25, 5.0)
                add_text_box(slide, "Wikimedia Commons", 9.0, 6.55, 3.1, .25, 8, False, MUTED, align=PP_ALIGN.RIGHT)
            add_number(slide, idx, accent)

        elif layout_no == 1:
            # CARDS — әр ой жеке карточка
            rect(slide, 0, 0, 13.333, 7.5, WHITE)
            add_text_box(slide, stitle, .7, .65, 11.6, .8, 28, True, INK)
            add_text_box(slide, "Негізгі ойлар", .72, 1.45, 4, .35, 12, True, accent)
            count = min(max(len(bullets), 1), 4)
            cols = 2 if count > 1 else 1
            rows = (count + cols - 1) // cols
            card_w = 5.7 if cols == 2 else 11.5
            for j in range(count):
                r, c = divmod(j, cols)
                x = .7 + c * 6.0; y = 2.0 + r * 2.05
                rect(slide, x, y, card_w, 1.65, LIGHT, True)
                rect(slide, x, y, .09, 1.65, accent)
                txt = bullets[j] if j < len(bullets) else "Тақырыптың негізгі мазмұны."
                add_text_box(slide, f"{j+1:02d}", x+.3, y+.25, .6, .4, 12, True, accent)
                add_text_box(slide, txt, x+.95, y+.22, card_w-1.25, 1.1, 16, False, INK)
            add_number(slide, idx, accent)

        elif layout_no == 2:
            # FULL IMAGE — сурет басым, төменде мәтіндік панель
            rect(slide, 0, 0, 13.333, 7.5, NAVY)
            if image:
                add_image_safely(slide, image["stream"], 0, 0, 13.333, 7.5)
                rect(slide, 0, 4.75, 13.333, 2.75, NAVY)
            else:
                rect(slide, 0, 0, 13.333, 7.5, NAVY)
            add_text_box(slide, stitle, .75, 5.05, 8.9, .75, 28, True, WHITE)
            short = "  •  ".join(bullets[:2]) if bullets else "Тақырып бойынша негізгі түсініктер мен маңызды тұжырымдар."
            add_text_box(slide, short, .78, 5.95, 11.3, .85, 15, False, (225, 230, 240))
            if image:
                add_text_box(slide, "Wikimedia Commons", 10.2, 6.95, 2.3, .25, 8, False, WHITE, align=PP_ALIGN.RIGHT)
            add_number(slide, idx, WHITE)

        else:
            # EDITORIAL — үлкен сан + мәтін + кішкентай сурет
            rect(slide, 0, 0, 13.333, 7.5, (249, 250, 253))
            add_text_box(slide, f"{idx:02d}", .7, .7, 2.0, 1.3, 48, True, accent)
            add_text_box(slide, stitle, 2.35, .85, 6.9, 1.0, 28, True, INK)
            rect(slide, .75, 2.0, 7.3, .04, accent)
            add_bullets(slide, bullets, .8, 2.45, 6.9, 3.7, 18)
            if image:
                rect(slide, 8.65, 1.35, 3.9, 4.95, WHITE, True)
                add_image_safely(slide, image["stream"], 8.82, 1.52, 3.56, 4.58)
                add_text_box(slide, "Wikimedia Commons", 9.05, 6.2, 3.0, .25, 8, False, MUTED, align=PP_ALIGN.RIGHT)
            add_number(slide, idx, accent)

    out_dir = Path("generated")
    out_dir.mkdir(exist_ok=True)
    path = out_dir / f"presentation_{order_id}.pptx"
    prs.save(path)
    return path


# =========================
# HTML
# =========================
def layout(title, body):
    return f"""<!doctype html>
<html lang="kk">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(title)}</title>

<style>
body {{
    margin: 0;
    background: #f5f7fb;
    font-family: Arial, sans-serif;
    color: #222;
}}

.box {{
    max-width: 650px;
    margin: 35px auto;
    padding: 28px;
    background: #fff;
    border-radius: 20px;
    box-shadow: 0 5px 25px rgba(0,0,0,.08);
}}

h1, h2 {{
    margin-top: 0;
}}

label {{
    display: block;
    margin: 14px 0 6px;
    font-weight: 600;
}}

input, select, textarea {{
    width: 100%;
    box-sizing: border-box;
    padding: 12px;
    border: 1px solid #ddd;
    border-radius: 10px;
    font-size: 16px;
}}

textarea {{
    min-height: 100px;
    resize: vertical;
}}

button, .button {{
    display: inline-block;
    border: 0;
    margin-top: 18px;
    padding: 13px 20px;
    background: #7c3aed;
    color: #fff;
    border-radius: 10px;
    font-size: 16px;
    text-decoration: none;
    cursor: pointer;
}}

.secondary {{
    background: #555;
}}

.price {{
    font-size: 24px;
    font-weight: bold;
    margin: 18px 0;
}}

.free {{
    color: #138a3d;
    font-weight: bold;
}}

.note {{
    background: #f2efff;
    padding: 12px;
    border-radius: 10px;
    margin: 15px 0;
}}

.error {{
    background: #ffecec;
    padding: 12px;
    border-radius: 10px;
    color: #a00;
}}

.row {{
    display: flex;
    gap: 10px;
}}

.row > * {{
    flex: 1;
}}

small {{
    color: #666;
}}
</style>
</head>

<body>
<div class="box">
{body}
</div>
</body>
</html>"""


def page(title, body, status=200):
    return layout(title, body), status


# =========================
# HOME
# =========================
@app.route("/")
def home():
    customer_id = customer_id_from_request()

    if not customer_id:
        customer_id = create_customer()
        ensure_customer(customer_id)

        resp = make_response(
            _home_html(customer_id)
        )

        resp.set_cookie(
            "tapsyrys_customer_id",
            customer_id,
            max_age=60 * 60 * 24 * 365,
            httponly=True,
            samesite="Lax",
            secure=True,
        )

        return resp

    ensure_customer(customer_id)

    return _home_html(customer_id)


def _home_html(customer_id):
    customer = get_customer(customer_id)

    presentation_free_available = not bool(
        customer["presentation_free_used"]
    )

    options = "".join(
        f'<option value="{html.escape(name)}">'
        f'{html.escape(name)}</option>'
        for name in SERVICES
    )

    free_text = (
        '<div class="note free">'
        '🎁 1 тегін презентация жасау мүмкіндігі бар!'
        '</div>'
        if presentation_free_available
        else
        '<div class="note">'
        '🎁 Тегін мүмкіндік қолданылды. '
        'Келесі презентациялар — 60 тг/слайд.'
        '</div>'
    )

    return layout(
        "Tapsyrys AI",
        f"""
<h1>Tapsyrys AI 🚀</h1>

<p>
Тақырыпты енгізіңіз — бірінші презентация тегін 🎁
</p>

{free_text}

<form method="post" action="/order">

<label>Қызмет</label>

<select
    name="service"
    id="service"
    onchange="updatePrice()"
    required
>
{options}
</select>

<label>Пән</label>

<input
    name="subject"
    placeholder="Мысалы: Философия"
>

<label>Тақырып</label>

<textarea
    name="topic"
    placeholder="Мысалы: Болмыс және таным"
    required
></textarea>

<label>Слайд саны</label>

<input
    name="quantity"
    id="quantity"
    type="number"
    min="1"
    max="{MAX_SLIDES}"
    value="10"
    oninput="updatePrice()"
>

<div class="row">

<div>
<label>Күні</label>
<input name="due_date" type="date">
</div>

<div>
<label>Уақыты</label>
<input name="due_time" type="time">
</div>

</div>

<div class="price" id="price">
Бағасы: 0 тг
</div>

<button type="submit">
Тапсырыс беру
</button>

</form>

<hr style="margin:25px 0">

<a
    class="button secondary"
    href="/canva/login"
>
Canva аккаунтын қосу
</a>

<p>
<small>
Суреттер автоматты түрде Wikimedia Commons-тен
тегін түрде ізделеді.
</small>
</p>

<script>
const services =
    {json.dumps(SERVICES, ensure_ascii=False)};

const presentationFreeAvailable =
    {str(presentation_free_available).lower()};

function updatePrice() {{
    const service =
        document.getElementById("service").value;

    const quantity =
        Math.max(
            1,
            parseInt(
                document.getElementById("quantity").value || "1"
            )
        );

    let price =
        services[service].price;

    if (services[service].type !== "fixed") {{
        price *= quantity;
    }}

    if (
        service === "Презентация"
        && presentationFreeAvailable
    ) {{
        document.getElementById("price").innerHTML =
            'Бағасы: <span class="free">'
            + '0 тг — бірінші презентация тегін 🎁'
            + '</span>';
    }} else {{
        document.getElementById("price").textContent =
            "Бағасы: " + price + " тг";
    }}
}}

updatePrice();
</script>
""",
    )


# =========================
# ORDER
# =========================
@app.route("/order", methods=["POST"])
def order():
    customer_id = customer_id_from_request()

    if not customer_id:
        return redirect("/")

    ensure_customer(customer_id)

    service = request.form.get(
        "service",
        ""
    )

    subject = request.form.get(
        "subject",
        ""
    ).strip()

    topic = request.form.get(
        "topic",
        ""
    ).strip()

    due_date = request.form.get(
        "due_date",
        ""
    )

    due_time = request.form.get(
        "due_time",
        ""
    )

    try:
        quantity = max(
            1,
            min(
                MAX_SLIDES,
                int(
                    request.form.get(
                        "quantity",
                        "1"
                    )
                ),
            ),
        )
    except ValueError:
        quantity = 1

    if service not in SERVICES or not topic:
        return page(
            "Қате",
            """
<h2>Мәлімет толық емес ❌</h2>
<a class="button" href="/">Қайту</a>
""",
            400,
        )

    customer = get_customer(customer_id)

    presentation_free_available = (
        service == "Презентация"
        and not bool(
            customer["presentation_free_used"]
        )
    )

    normal_price = calculate_price(
        service,
        quantity
    )

    price = (
        0
        if presentation_free_available
        else normal_price
    )

    with db() as conn:
        cur = conn.execute(
            """
            INSERT INTO orders
            (
                customer_id,
                service,
                subject,
                topic,
                quantity,
                due_date,
                due_time,
                price,
                payment_status,
                generation_status,
                created_at
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                customer_id,
                service,
                subject,
                topic,
                quantity,
                due_date,
                due_time,
                price,
                (
                    "free"
                    if presentation_free_available
                    else "pending"
                ),
                (
                    "unlocked"
                    if presentation_free_available
                    else "locked"
                ),
                time.time(),
            ),
        )

        order_id = cur.lastrowid

        if presentation_free_available:
            conn.execute(
                """
                UPDATE customers
                SET presentation_free_used = 1
                WHERE customer_id = ?
                """,
                (customer_id,),
            )

    if presentation_free_available:
        return redirect(
            f"/order/{order_id}/generation"
        )

    return page(
        "Төлем",
        f"""
<h2>Тапсырыс #{order_id}</h2>

<p>
<b>Қызмет:</b>
{html.escape(service)}
</p>

<p>
<b>Тақырып:</b>
{html.escape(topic)}
</p>

<p>
<b>Слайд саны:</b>
{quantity}
</p>

<div class="price">
Төлем: {normal_price} тг
</div>

<div class="note">
Бұл әзірге <b>тесттік төлем</b>.
Нақты ақша алынбайды.
</div>

<form
    method="post"
    action="/order/{order_id}/test-payment"
>
<button type="submit">
Тест төлемін растау
</button>
</form>

<a class="button secondary" href="/">
Басты бет
</a>
""",
    )


@app.route(
    "/order/<int:order_id>/test-payment",
    methods=["POST"]
)
def test_payment(order_id):
    customer_id = customer_id_from_request()

    with db() as conn:
        order = conn.execute(
            """
            SELECT *
            FROM orders
            WHERE id = ?
            AND customer_id = ?
            """,
            (
                order_id,
                customer_id,
            ),
        ).fetchone()

        if not order:
            return page(
                "Қате",
                "<h2>Тапсырыс табылмады ❌</h2>",
                404,
            )

        conn.execute(
            """
            UPDATE orders
            SET payment_status = 'paid',
                generation_status = 'unlocked'
            WHERE id = ?
            """,
            (order_id,),
        )

    return redirect(
        f"/order/{order_id}/generation"
    )


# =========================
# GENERATION PAGE
# =========================
@app.route("/order/<int:order_id>/generation")
def generation(order_id):
    customer_id = customer_id_from_request()

    with db() as conn:
        order = conn.execute(
            """
            SELECT *
            FROM orders
            WHERE id = ?
            AND customer_id = ?
            """,
            (
                order_id,
                customer_id,
            ),
        ).fetchone()

    if not order:
        return page(
            "Қате",
            "<h2>Тапсырыс табылмады ❌</h2>",
            404,
        )

    if order["generation_status"] != "unlocked":
        return page(
            "Құлыпталған",
            """
<h2>Генерация жабық 🔒</h2>
<p>Алдымен төлемді растаңыз.</p>
""",
            403,
        )

    return page(
        "Генерация",
        f"""
<h2>Презентация жасауға дайын ✨</h2>

<p>
<b>Тапсырыс:</b>
#{order["id"]}
</p>

<p>
<b>Пән:</b>
{html.escape(order["subject"] or "Көрсетілмеген")}
</p>

<p>
<b>Тақырып:</b>
{html.escape(order["topic"])}
</p>

<p>
<b>Слайд:</b>
{order["quantity"]}
</p>

<div class="note">
Әр слайдқа тақырыпқа сәйкес
Wikimedia Commons-тен тегін сурет ізделеді.
OpenAI API қажет емес.
</div>

<form
    method="post"
    action="/order/{order_id}/generate"
>
<button type="submit">
Генерация жасау ✨
</button>
</form>

<a class="button secondary" href="/">
Басты бет
</a>
""",
    )


# =========================
# GENERATE PPTX
# =========================
@app.route(
    "/order/<int:order_id>/generate",
    methods=["POST"]
)
def generate(order_id):
    customer_id = customer_id_from_request()

    with db() as conn:
        order = conn.execute(
            """
            SELECT *
            FROM orders
            WHERE id = ?
            AND customer_id = ?
            """,
            (
                order_id,
                customer_id,
            ),
        ).fetchone()

    if (
        not order
        or order["generation_status"] != "unlocked"
    ):
        return page(
            "Қате",
            "<h2>Генерацияға рұқсат жоқ ❌</h2>",
            403,
        )

    try:
        with db() as conn:
            conn.execute(
                """
                UPDATE orders
                SET generation_status = 'generating'
                WHERE id = ?
                """,
                (order_id,),
            )

        slides_data = make_local_slides(
            order["topic"],
            order["subject"],
            order["quantity"],
        )

        make_pptx(
            slides_data,
            order_id,
        )

        with db() as conn:
            conn.execute(
                """
                UPDATE orders
                SET generation_status = 'completed'
                WHERE id = ?
                """,
                (order_id,),
            )

        return page(
            "Дайын презентация",
            f"""
<h2>Презентация дайын! 🎉</h2>

<p>
<b>Тапсырыс:</b>
#{order_id}
</p>

<p>
<b>Тақырып:</b>
{html.escape(order["topic"])}
</p>

<div class="note">
PPTX жасалды.
Суреттер Wikimedia Commons-тен
автоматты түрде ізделді.
</div>

<a
    class="button"
    href="/download/{order_id}"
>
PPTX жүктеу 📥
</a>

<br>

<a
    class="button secondary"
    href="/"
>
Басты бет
</a>
""",
        )

    except Exception as e:
        with db() as conn:
            conn.execute(
                """
                UPDATE orders
                SET generation_status = 'unlocked'
                WHERE id = ?
                """,
                (order_id,),
            )

        return page(
            "Генерация қатесі",
            f"""
<h2>Презентация жасау кезінде қате ❌</h2>

<pre>{html.escape(str(e))}</pre>

<a class="button" href="/">
Басты бет
</a>
""",
            500,
        )


# =========================
# DOWNLOAD
# =========================
@app.route("/download/<int:order_id>")
def download_presentation(order_id):
    customer_id = customer_id_from_request()

    with db() as conn:
        order = conn.execute(
            """
            SELECT *
            FROM orders
            WHERE id = ?
            AND customer_id = ?
            """,
            (
                order_id,
                customer_id,
            ),
        ).fetchone()

    path = (
        Path("generated")
        / f"presentation_{order_id}.pptx"
    )

    if (
        not order
        or order["service"] != "Презентация"
        or order["generation_status"] != "completed"
        or not path.exists()
    ):
        return page(
            "Файл жоқ",
            "<h2>Файл табылмады ❌</h2>",
            404,
        )

    return send_file(
        path,
        as_attachment=True,
        download_name=(
            f"Tapsyrys_presentation_{order_id}.pptx"
        ),
    )


# =========================
# CANVA LOGIN
# =========================
@app.route("/canva/login")
def canva_login():
    if not CLIENT_ID:
        return page(
            "Қате",
            """
<h2>Қате ❌</h2>
<p>CANVA_CLIENT_ID табылмады.</p>
""",
            500,
        )

    state = secrets.token_urlsafe(32)

    code_verifier, code_challenge = make_pkce()

    save_state(
        state,
        code_verifier,
    )

    params = {
        "client_id": CLIENT_ID,
        "redirect_uri": REDIRECT_URI,
        "response_type": "code",
        "scope": SCOPES,
        "state": state,
        "code_challenge": code_challenge,
        "code_challenge_method": "S256",
    }

    return redirect(
        CANVA_AUTHORIZE_URL
        + "?"
        + urllib.parse.urlencode(params)
    )


# =========================
# CANVA CALLBACK
# =========================
@app.route("/canva/callback")
def canva_callback():
    error = request.args.get("error")

    if error:
        description = request.args.get(
            "error_description",
            "Canva OAuth қатесі",
        )

        return page(
            "OAuth қатесі",
            f"""
<h2>Canva OAuth қатесі ❌</h2>

<p>
<b>Error:</b>
{html.escape(error)}
</p>

<p>
<b>Description:</b>
{html.escape(description)}
</p>

<a href="/">
Басты бетке қайту
</a>
""",
            400,
        )

    code = request.args.get("code")
    state = request.args.get("state")

    if not code:
        return page(
            "Код жоқ",
            """
<h2>Authorization code жоқ ❌</h2>
<a href="/">Басты бетке қайту</a>
""",
            400,
        )

    code_verifier = (
        pop_state(state)
        if state
        else None
    )

    if not code_verifier:
        return page(
            "State қатесі",
            """
<h2>OAuth State қатесі ❌</h2>
<a href="/">Басты бетке қайту</a>
""",
            400,
        )

    if not CLIENT_ID or not CLIENT_SECRET:
        return page(
            "Қате",
            """
<h2>
CANVA_CLIENT_ID немесе
CANVA_CLIENT_SECRET жоқ ❌
</h2>
""",
            500,
        )

    try:
        credentials = (
            f"{CLIENT_ID}:{CLIENT_SECRET}"
        )

        encoded_credentials = (
            base64.b64encode(
                credentials.encode("utf-8")
            ).decode("utf-8")
        )

        data = urllib.parse.urlencode(
            {
                "grant_type":
                    "authorization_code",
                "code":
                    code,
                "redirect_uri":
                    REDIRECT_URI,
                "code_verifier":
                    code_verifier,
            }
        ).encode("utf-8")

        req = urllib.request.Request(
            CANVA_TOKEN_URL,
            data=data,
            method="POST",
        )

        req.add_header(
            "Authorization",
            "Basic "
            + encoded_credentials,
        )

        req.add_header(
            "Content-Type",
            "application/x-www-form-urlencoded",
        )

        req.add_header(
            "Accept",
            "application/json",
        )

        with urllib.request.urlopen(
            req,
            timeout=30,
        ) as response:
            result = (
                response
                .read()
                .decode("utf-8")
            )

        token_data = json.loads(result)

        if not token_data.get(
            "access_token"
        ):
            safe = {
                k: v
                for k, v in token_data.items()
                if k in (
                    "error",
                    "error_description",
                )
            }

            return page(
                "Токен алынбады",
                f"""
<pre>
{html.escape(
    json.dumps(
        safe,
        indent=2,
        ensure_ascii=False
    )
)}
</pre>
""",
                400,
            )

        save_tokens(token_data)

        return page(
            "Canva Connected",
            """
<h1>Canva сәтті қосылды! ✅</h1>

<p>
Canva аккаунты Tapsyrys AI жүйесіне қосылды.
</p>

<a class="button" href="/">
Басты бетке қайту
</a>
""",
        )

    except urllib.error.HTTPError as e:
        error_body = e.read().decode(
            "utf-8",
            errors="ignore",
        )

        return page(
            "Token қатесі",
            f"""
<h2>Canva Token қатесі ❌</h2>

<p>
HTTP: {e.code}
</p>

<pre>
{html.escape(error_body)}
</pre>

<a href="/">
Басты бетке қайту
</a>
""",
            400,
        )

    except Exception as e:
        return page(
            "Сервер қатесі",
            f"""
<h2>Сервер қатесі ❌</h2>

<pre>
{html.escape(str(e))}
</pre>
""",
            500,
        )


# =========================
# HEALTH
# =========================
@app.route("/health")
def health():
    return {
        "status": "ok",
        "service": "Tapsyrys AI",
        "images": "Wikimedia Commons",
        "openai": False,
    }


# =========================
# PKCE
# =========================
def make_pkce():
    code_verifier = secrets.token_urlsafe(64)

    digest = hashlib.sha256(
        code_verifier.encode("ascii")
    ).digest()

    code_challenge = (
        base64.urlsafe_b64encode(digest)
        .rstrip(b"=")
        .decode("ascii")
    )

    return (
        code_verifier,
        code_challenge,
    )


# =========================
# START
# =========================
if __name__ == "__main__":
    port = int(
        os.environ.get(
            "PORT",
            10000,
        )
    )

    app.run(
        host="0.0.0.0",
        port=port,
    )
